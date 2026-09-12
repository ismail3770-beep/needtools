"""
Office / image-format conversion endpoints for the NeedTools PDF service.

Mounted on the compress service via `app.include_router(office_router)`.

Requires (already present in the Docker image):
  - libreoffice  : office -> pdf
  - PyMuPDF      : pdf rendering
  - Pillow       : TIFF encode/decode
  - python-pptx  : pptx generation
"""

import asyncio
import io
import os
import subprocess
import uuid

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

import pymupdf as fitz

router = APIRouter()

TMP_DIR = "/tmp" if os.name != "nt" else "."

# LibreOffice can hang on malformed documents; always bound it.
SOFFICE_TIMEOUT = 180
MAX_UPLOAD_BYTES = 30 * 1024 * 1024

# Extensions LibreOffice can reliably turn into PDF.
OFFICE_TO_PDF_EXTS = {
    ".ppt", ".pptx", ".odp",
    ".xls", ".xlsx", ".ods", ".csv",
    ".doc", ".docx", ".odt", ".rtf",
    ".txt",
}


def _tmp(suffix: str) -> str:
    return os.path.join(TMP_DIR, f"{uuid.uuid4()}{suffix}")


def _ext_of(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def _safe_stem(filename: str) -> str:
    stem = os.path.splitext(os.path.basename(filename or "document"))[0]
    return (stem or "document").replace('"', "")


async def _read_upload(file: UploadFile) -> bytes:
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=400,
            detail=f"File too large. Max {MAX_UPLOAD_BYTES // (1024 * 1024)}MB.",
        )
    return content


def _cleanup(*paths: str) -> None:
    for p in paths:
        try:
            if p and os.path.exists(p):
                os.remove(p)
        except OSError:
            pass


# ---------------------------------------------------------------------------
# Office -> PDF (LibreOffice headless)
# ---------------------------------------------------------------------------
def _do_office_to_pdf(input_path: str) -> bytes:
    outdir = os.path.dirname(input_path)
    try:
        subprocess.run(
            [
                "libreoffice", "--headless", "--nologo", "--nofirststartwizard",
                "--convert-to", "pdf", input_path, "--outdir", outdir,
            ],
            check=True,
            timeout=SOFFICE_TIMEOUT,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError("Conversion timed out. The document may be too complex.")
    except subprocess.CalledProcessError as exc:
        detail = (exc.stderr or b"").decode("utf-8", "ignore").strip()
        raise RuntimeError(detail or "LibreOffice failed to convert the document.")

    produced = os.path.join(
        outdir, os.path.splitext(os.path.basename(input_path))[0] + ".pdf"
    )
    if not os.path.exists(produced):
        raise RuntimeError("LibreOffice produced no output file.")

    try:
        with open(produced, "rb") as f:
            return f.read()
    finally:
        _cleanup(produced)


@router.post("/convert/office-to-pdf")
async def office_to_pdf_endpoint(file: UploadFile = File(...)):
    """Convert PowerPoint, Excel, text and OpenDocument files to PDF."""
    ext = _ext_of(file.filename)
    if ext not in OFFICE_TO_PDF_EXTS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{ext or 'unknown'}'.",
        )

    content = await _read_upload(file)
    input_path = _tmp("_office_input" + ext)

    try:
        with open(input_path, "wb") as f:
            f.write(content)

        loop = asyncio.get_running_loop()
        result_bytes = await loop.run_in_executor(None, _do_office_to_pdf, input_path)

        stem = _safe_stem(file.filename)
        return StreamingResponse(
            io.BytesIO(result_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{stem}.pdf"',
                "Content-Length": str(len(result_bytes)),
            },
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception as exc:
        print(f"[office_to_pdf] Error: {exc}")
        raise HTTPException(status_code=500, detail=str(exc))
    finally:
        _cleanup(input_path)


# ---------------------------------------------------------------------------
# PDF -> PowerPoint (one slide per page, page rendered as a full-bleed image)
# ---------------------------------------------------------------------------
def _do_pdf_to_ppt(pdf_bytes: bytes, dpi: int) -> bytes:
    from pptx import Presentation
    from pptx.util import Emu

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    if doc.page_count == 0:
        doc.close()
        raise RuntimeError("The PDF has no pages.")

    zoom = max(0.5, min(4.0, dpi / 72.0))
    matrix = fitz.Matrix(zoom, zoom)

    prs = Presentation()
    blank_layout = prs.slide_layouts[6]  # 6 = blank

    # Size the deck from the first page so slides match the document aspect ratio.
    first = doc.load_page(0)
    prs.slide_width = Emu(int(first.rect.width * 12700))
    prs.slide_height = Emu(int(first.rect.height * 12700))

    try:
        for index in range(doc.page_count):
            page = doc.load_page(index)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            image_stream = io.BytesIO(pix.tobytes("png"))

            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                image_stream,
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )
    finally:
        doc.close()

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


@router.post("/convert/pdf-to-ppt")
async def pdf_to_ppt_endpoint(
    file: UploadFile = File(...),
    dpi: int = Form(default=150),
):
    """Convert each PDF page into a PowerPoint slide image."""
    if _ext_of(file.filename) != ".pdf" and file.content_type != "application/pdf":
        raise HTTPException(status_code=400, detail="Only PDF files are accepted.")

    content = await _read_upload(file)
    dpi = max(72, min(300, dpi))

    try:
        loop = asyncio.get_running_loop()
        result_bytes = await loop.run_in_executor(None, _do_pdf_to_ppt, content, dpi)

        stem = _safe_stem(file.filename)
        return StreamingResponse(
            io.BytesIO(result_bytes),
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={
                "Content-Disposition": f'attachment; filename="{stem}.pptx"',
                "Content-Length": str(len(result_bytes)),
            },
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception as exc:
        print(f"[pdf_to_ppt] Error: {exc}")
        raise HTTPException(status_code=500, detail=str(exc))


# ---------------------------------------------------------------------------
# PDF -> multi-page TIFF
# ---------------------------------------------------------------------------
def _do_pdf_to_tiff(pdf_bytes: bytes, dpi: int) -> bytes:
    import PIL.Image

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    if doc.page_count == 0:
        doc.close()
        raise RuntimeError("The PDF has no pages.")

    zoom = max(0.5, min(4.0, dpi / 72.0))
    matrix = fitz.Matrix(zoom, zoom)

    frames = []
    try:
        for index in range(doc.page_count):
            page = doc.load_page(index)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            frames.append(
                PIL.Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
            )
    finally:
        doc.close()

    output = io.BytesIO()
    frames[0].save(
        output,
        format="TIFF",
        save_all=True,
        append_images=frames[1:],
        compression="tiff_deflate",
        dpi=(dpi, dpi),
    )
    return output.getvalue()


@router.post("/convert/pdf-to-tiff")
async def pdf_to_tiff_endpoint(
    file: UploadFile = File(...),
    dpi: int = Form(default=150),
):
    """Convert a PDF into a single multi-page TIFF."""
    if _ext_of(file.filename) != ".pdf" and file.content_type != "application/pdf":
        raise HTTPException(status_code=400, detail="Only PDF files are accepted.")

    content = await _read_upload(file)
    dpi = max(72, min(300, dpi))

    try:
        loop = asyncio.get_running_loop()
        result_bytes = await loop.run_in_executor(None, _do_pdf_to_tiff, content, dpi)

        stem = _safe_stem(file.filename)
        return StreamingResponse(
            io.BytesIO(result_bytes),
            media_type="image/tiff",
            headers={
                "Content-Disposition": f'attachment; filename="{stem}.tiff"',
                "Content-Length": str(len(result_bytes)),
            },
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception as exc:
        print(f"[pdf_to_tiff] Error: {exc}")
        raise HTTPException(status_code=500, detail=str(exc))


# ---------------------------------------------------------------------------
# TIFF -> PDF (every frame of a multi-page TIFF becomes one PDF page)
# ---------------------------------------------------------------------------
def _do_tiff_to_pdf(tiff_bytes: bytes) -> bytes:
    import PIL.Image
    from PIL import ImageSequence

    source = PIL.Image.open(io.BytesIO(tiff_bytes))
    frames = [frame.convert("RGB") for frame in ImageSequence.Iterator(source)]
    if not frames:
        raise RuntimeError("No image frames found in the TIFF file.")

    output = io.BytesIO()
    frames[0].save(
        output,
        format="PDF",
        save_all=True,
        append_images=frames[1:],
        resolution=96.0,
    )
    return output.getvalue()


@router.post("/convert/tiff-to-pdf")
async def tiff_to_pdf_endpoint(file: UploadFile = File(...)):
    """Convert a single-page or multi-page TIFF into a PDF."""
    if _ext_of(file.filename) not in {".tif", ".tiff"}:
        raise HTTPException(status_code=400, detail="Only TIFF files are accepted.")

    content = await _read_upload(file)

    try:
        loop = asyncio.get_running_loop()
        result_bytes = await loop.run_in_executor(None, _do_tiff_to_pdf, content)

        stem = _safe_stem(file.filename)
        return StreamingResponse(
            io.BytesIO(result_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{stem}.pdf"',
                "Content-Length": str(len(result_bytes)),
            },
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except Exception as exc:
        print(f"[tiff_to_pdf] Error: {exc}")
        raise HTTPException(status_code=500, detail=str(exc))
